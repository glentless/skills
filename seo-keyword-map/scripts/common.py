"""Shared helpers: project config, keyword-map workbook access, DataForSEO client.

Python 3.9+. Needs openpyxl for anything that touches the workbook:
    pip install openpyxl
"""
from __future__ import annotations

import base64
import csv
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import sys
import urllib.error
import urllib.request
from pathlib import Path

for _s in (sys.stdout, sys.stderr):  # Windows consoles default to cp1252 and mangle — and ·
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

SKILL_DIR = Path(__file__).resolve().parent.parent
CONFIG_NAME = "seo-project.json"
TODAY = dt.date.today().isoformat()

# Standard keyword-map.xlsx schema. Config "map.columns" can rename any of these.
MAP_COLUMNS = [
    "Service line", "Keyword", "Volume", "Difficulty", "Intent", "Cluster",
    "Current rank", "Best competitor", "Target page", "Page exists", "Status",
    "Priority", "Action", "Role", "Cluster volume", "Page type",
]
ROLES = ("Primary", "Secondary", "Support")
PAGE_TYPES = ("hub", "service", "location", "post")
STATUSES = ("Mapped to page", "No page yet", "Unmapped", "Excluded")
TRIAGE_COLUMNS = ["Keyword", "Volume", "Disposition", "Where it went, or why not"]
DISPOSITIONS = ("Added to Keyword map", "Folded into a kept term", "Removed")


def die(msg: str, code: int = 2):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(code)


# ---------------------------------------------------------------- config

def find_project(start: Path | None = None) -> Path | None:
    """Walk up from cwd to the folder holding seo-project.json."""
    p = (start or Path.cwd()).resolve()
    for d in [p, *p.parents]:
        if (d / CONFIG_NAME).exists():
            return d
    return None


def load_config(required: bool = True) -> tuple[dict, Path | None]:
    root = find_project()
    if not root:
        if required:
            die(f"No {CONFIG_NAME} found here or in any parent folder. "
                "Run: python <skill>/scripts/init_project.py --domain example.com")
        return {}, None
    try:
        cfg = json.loads((root / CONFIG_NAME).read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        die(f"{CONFIG_NAME} is not valid JSON: {e}")
    cfg.setdefault("map", {}).setdefault("path", "keyword-map.xlsx")
    cfg.setdefault("data_dir", "seo-data")
    cfg.setdefault("min_volume", 20)
    cfg.setdefault("budget", {}).setdefault("max_usd_per_run", 1.00)
    return cfg, root


def data_dir(cfg: dict, root: Path) -> Path:
    d = root / cfg["data_dir"]
    d.mkdir(parents=True, exist_ok=True)
    return d


def col(cfg: dict, name: str) -> str:
    """Header text used in this project's workbook for a standard column."""
    return cfg.get("map", {}).get("columns", {}).get(name, name)


# ---------------------------------------------------------------- workbook

def map_path(cfg: dict, root: Path) -> Path:
    return root / cfg["map"]["path"]


def lock_file(path: Path) -> Path:
    return path.with_name("~$" + path.name)


def check_not_open(path: Path):
    """Excel holds ~$name while the file is open, and a save then fails or is lost."""
    if lock_file(path).exists():
        die(f"{path.name} is open in Excel (found {lock_file(path).name}). "
            "Close it, then run this again.")


def backup(path: Path) -> Path:
    stem = f"{path.stem}-backup-{TODAY}"
    b = path.with_name(stem + path.suffix)
    n = 2
    while b.exists():
        b = path.with_name(f"{stem}-{n}{path.suffix}")
        n += 1
    shutil.copyfile(path, b)
    return b


def open_map(cfg: dict, root: Path, read_only: bool = False):
    try:
        import openpyxl
    except ImportError:
        die("openpyxl is missing. Install it with: pip install openpyxl")
    p = map_path(cfg, root)
    if not p.exists():
        die(f"Keyword map not found at {p}. Run init_project.py, or set map.path in {CONFIG_NAME}.")
    wb = openpyxl.load_workbook(p, read_only=read_only, data_only=read_only)
    sheet = cfg["map"].get("sheet", "Keyword map")
    if sheet not in wb.sheetnames:
        die(f"Sheet '{sheet}' not found in {p.name}. Available: {', '.join(wb.sheetnames)}")
    return wb, wb[sheet]


def header_index(ws, cfg: dict, need=("Keyword", "Volume", "Target page", "Page exists",
                                      "Status", "Priority", "Role")) -> dict:
    """Map standard column name -> 0-based index. Missing optional columns are absent."""
    head = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {}
    for std in MAP_COLUMNS:
        h = col(cfg, std)
        if h in head:
            idx[std] = head.index(h)
    missing = [n for n in need if n not in idx]
    if missing:
        die(f"Column(s) {missing} not found in '{ws.title}'. "
            f"Available: {', '.join(str(h) for h in head if h)}. "
            f"Rename them in {CONFIG_NAME} under map.columns if this project uses other headers.")
    return idx


def read_map(cfg: dict, root: Path) -> list[dict]:
    """Rows of the keyword map as dicts keyed by standard column names."""
    wb, ws = open_map(cfg, root, read_only=True)
    idx = header_index(ws, cfg)
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or idx["Keyword"] >= len(r) or not r[idx["Keyword"]]:
            continue
        d = {k: (r[i] if i < len(r) else None) for k, i in idx.items()}
        d["Keyword"] = str(d["Keyword"]).strip()
        d["Volume"] = as_int(d.get("Volume"))
        rows.append(d)
    wb.close()
    return rows


def as_int(v):
    if v is None or v == "":
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def pages_from_rows(rows: list[dict]) -> dict:
    """Group map rows by target page."""
    pages: dict = {}
    for r in rows:
        p = r.get("Target page")
        if not p or r.get("Status") in ("Excluded", "Unmapped"):
            continue
        pg = pages.setdefault(p, {"url": p, "keywords": [], "volume": 0, "exists": None,
                                  "priority": None, "type": None, "primary": None,
                                  "lines": {}})
        pg["keywords"].append(r)
        pg["volume"] += r.get("Volume") or 0
        ex = str(r.get("Page exists") or "").strip().lower()
        pg["exists"] = ex in ("yes", "true", "1") if pg["exists"] is None else pg["exists"]
        pg["priority"] = pg["priority"] or as_int(r.get("Priority"))
        pg["type"] = pg["type"] or (r.get("Page type") or None)
        if r.get("Role") == "Primary":
            pg["primary"] = r["Keyword"]
        line = r.get("Service line") or ""
        pg["lines"][line] = pg["lines"].get(line, 0) + 1
    for pg in pages.values():
        pg["line"] = max(pg["lines"], key=pg["lines"].get) if pg["lines"] else ""
    return pages


EVERY_DAYS = {"daily": 1, "weekly": 7, "biweekly": 14, "monthly": 31, "quarterly": 92}


def next_rank_due(last: str | None, ranks_cfg: dict) -> str | None:
    """Next due date for a scheduled rank check; None when checks are manual.
    Monthly/quarterly with a day: the next such day after the last check."""
    every = ranks_cfg.get("frequency")
    if not every:
        return None
    if not last:
        return TODAY
    d = dt.date.fromisoformat(last)
    day = ranks_cfg.get("day")
    if every in ("monthly", "quarterly") and day:
        step = 1 if every == "monthly" else 3
        y, m = d.year, d.month
        cand = dt.date(y, m, min(int(day), 28))
        while cand <= d:
            m += step
            y, m = y + (m - 1) // 12, (m - 1) % 12 + 1
            cand = dt.date(y, m, min(int(day), 28))
        return cand.isoformat()
    return (d + dt.timedelta(days=EVERY_DAYS[every])).isoformat()


def resolve_page(arg: str, urls) -> str | None:
    """Match a page argument to a target URL. Accepts 'slug', '/slug/', and the
    'C:/Program Files/Git/slug/' that Git Bash on Windows turns '/slug/' into."""
    a = (arg or "").replace("\\", "/").strip()
    want = "/" + a.strip("/") + "/"
    if want in urls:
        return want
    hits = [u for u in urls if want.endswith(u) and u != "/"]
    return max(hits, key=len) if hits else None


def norm_kw(k: str) -> str:
    return re.sub(r"\s+", " ", str(k).strip().lower())


# ---------------------------------------------------------------- snapshots

def snapshot_dir(cfg: dict, root: Path) -> Path:
    d = data_dir(cfg, root) / "serp"
    d.mkdir(parents=True, exist_ok=True)
    return d


def snapshots(cfg: dict, root: Path) -> list[Path]:
    """SERP snapshots, oldest first. A snapshot is a JSON file with a 'keywords' object."""
    out = []
    for p in sorted(snapshot_dir(cfg, root).glob("*.json")):
        if p.name.startswith("_"):
            continue
        try:
            with p.open(encoding="utf-8") as f:
                head = f.read(4096)
            if '"keywords"' in head or '"_note"' in head:
                out.append(p)
        except OSError:
            pass
    return out


def load_snapshot(path: Path) -> dict:
    d = json.loads(path.read_text(encoding="utf-8"))
    return {norm_kw(k): v for k, v in d.get("keywords", {}).items()}


def latest_serp(cfg: dict, root: Path) -> dict:
    """Merge every snapshot; newer data for a keyword replaces older."""
    merged = {}
    for p in snapshots(cfg, root):
        for k, v in load_snapshot(p).items():
            v = dict(v)
            v.setdefault("_snapshot", p.stem)
            merged[k] = v
    return merged


def serp_file(kw: str) -> str:
    """Raw-result file name for a keyword. The hash keeps 'as built' and 'as-built' apart."""
    base = re.sub(r"[^a-z0-9]+", "-", kw.lower()).strip("-")[:70]
    return f"{base}-{hashlib.md5(kw.lower().encode()).hexdigest()[:6]}"


def raw_titles(cfg: dict, root: Path, kw: str, date: str) -> list[str]:
    """Page-one organic titles for a keyword, from every pass of one pull's raw results."""
    out = []
    for f in sorted((snapshot_dir(cfg, root) / "raw" / date).glob(f"*/{serp_file(kw)}.json")):
        try:
            t = json.loads(f.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        for o in (t.get("organic") or [])[:10]:
            if o.get("title") and o["title"] not in out:
                out.append(o["title"])
    return out


# ---------------------------------------------------------------- DataForSEO

CRED_FILES = [  # first one found wins; none of these live inside a project or the skill
    Path.home() / ".dataforseo.env",
    Path.home() / ".config" / "dataforseo" / ".env",
    Path.home() / ".claude" / "skills" / "seo" / ".env",
]


def credentials() -> tuple[str, str, str] | None:
    """(login, password, where-from). Environment variables beat files."""
    lg, pw = os.environ.get("DATAFORSEO_LOGIN"), os.environ.get("DATAFORSEO_PASSWORD")
    if lg and pw:
        return lg, pw, "environment variables"
    for f in CRED_FILES:
        if f.exists():
            env = {}
            for line in f.read_text(encoding="utf-8").splitlines():
                if "=" in line and not line.strip().startswith("#"):
                    k, v = line.split("=", 1)
                    env[k.strip()] = v.strip().strip('"').strip("'")
            if env.get("DATAFORSEO_LOGIN") and env.get("DATAFORSEO_PASSWORD"):
                return env["DATAFORSEO_LOGIN"], env["DATAFORSEO_PASSWORD"], str(f)
    return None


def warn_stray_credentials(root: Path | None):
    """A credentials file inside a project or the skill would get committed."""
    for base in filter(None, [root, SKILL_DIR]):
        for f in list(base.glob(".env*")) + list(base.glob("*.env")):
            try:
                if "DATAFORSEO_PASSWORD" in f.read_text(encoding="utf-8", errors="ignore"):
                    print(f"WARNING: {f} holds DataForSEO credentials. Move it to "
                          f"{CRED_FILES[0]} so it is never committed.", file=sys.stderr)
            except OSError:
                pass


class DataForSEO:
    """Thin client that enforces a spending cap and logs every paid call to the project ledger."""

    BASE = "https://api.dataforseo.com/v3/"

    def __init__(self, cfg: dict, root: Path, max_usd: float | None = None, label: str = ""):
        cred = credentials()
        if not cred:
            die("DataForSEO credentials not found. Put DATAFORSEO_LOGIN and DATAFORSEO_PASSWORD "
                f"in {CRED_FILES[0]} (see reference/tutorial.md, setup step 2). "
                "Never put them inside a project folder.")
        warn_stray_credentials(root)
        self.auth = "Basic " + base64.b64encode(f"{cred[0]}:{cred[1]}".encode()).decode()
        self.cap = float(max_usd if max_usd is not None else cfg["budget"]["max_usd_per_run"])
        self.spent = 0.0
        self.label = label
        self.ledger = data_dir(cfg, root) / "dataforseo-ledger.csv"

    def _req(self, path: str, body=None) -> dict:
        req = urllib.request.Request(
            self.BASE + path,
            data=None if body is None else json.dumps(body).encode(),
            headers={"Authorization": self.auth, "Content-Type": "application/json"},
            method="GET" if body is None else "POST")
        for attempt in range(3):  # transient 5xx and timeouts happen; 3 tries is enough
            try:
                with urllib.request.urlopen(req, timeout=180) as r:
                    d = json.load(r)
                break
            except urllib.error.HTTPError as e:
                if e.code == 401:
                    die("DataForSEO rejected the login (401). Check the credentials file.")
                if e.code == 402:
                    die("DataForSEO account is out of credit (402). Top it up first.")
                if attempt == 2 or e.code < 500:
                    die(f"DataForSEO HTTP {e.code} on {path}: {e.read()[:300]!r}")
            except (urllib.error.URLError, TimeoutError) as e:
                if attempt == 2:
                    die(f"Network error calling DataForSEO {path}: {e}")
        if d.get("status_code") != 20000:
            die(f"DataForSEO error on {path}: {d.get('status_code')} {d.get('status_message')}")
        # task_get / tasks_ready repeat the cost of the task they return; only the post is charged
        cost = 0.0 if ("task_get" in path or "tasks_ready" in path) else float(d.get("cost") or 0)
        if cost:
            self.spent += cost
            new = not self.ledger.exists()
            with self.ledger.open("a", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                if new:
                    w.writerow(["date", "time", "endpoint", "label", "cost_usd"])
                w.writerow([TODAY, dt.datetime.now().strftime("%H:%M:%S"), path.split("/live")[0],
                            self.label, f"{cost:.5f}"])
        return d

    def preflight(self, estimate: float):
        """Stop before spending if the estimate breaks the cap or the balance."""
        if estimate > self.cap:
            die(f"Estimated cost ${estimate:.2f} is over the per-run cap of ${self.cap:.2f}. "
                "Narrow the request, or pass --max-cost with a higher figure if this is intended.")
        bal = self.balance()
        if bal is not None and bal < estimate:
            die(f"Estimated cost ${estimate:.2f} but the account balance is ${bal:.2f}.")
        print(f"Estimated cost: ${estimate:.3f} (cap ${self.cap:.2f}, balance ${bal:.2f})")

    def balance(self) -> float | None:
        try:
            d = self._req("appendix/user_data")
            return float(d["tasks"][0]["result"][0]["money"]["balance"])
        except Exception:
            return None

    def get(self, path: str) -> dict:
        return self._req(path)

    def post(self, path: str, tasks: list) -> dict:
        if self.spent > self.cap:
            die(f"Stopped: spent ${self.spent:.2f}, over the per-run cap of ${self.cap:.2f}.")
        return self._req(path, tasks)

    def done(self):
        print(f"DataForSEO spend this run: ${self.spent:.3f} (logged to {self.ledger.name})")


def ledger_total(cfg: dict, root: Path) -> float:
    f = data_dir(cfg, root) / "dataforseo-ledger.csv"
    if not f.exists():
        return 0.0
    with f.open(encoding="utf-8") as fh:
        return sum(float(r.get("cost_usd") or 0) for r in csv.DictReader(fh))
